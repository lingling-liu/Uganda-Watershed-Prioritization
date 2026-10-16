"""Read-only source extraction and reproducible spatial/scoring validation."""
from pathlib import Path
import sys,json,zipfile,collections,math,bisect,hashlib
sys.path.insert(0,str(Path(__file__).parent/'pydeps'))
import openpyxl,shapefile
from pyproj import CRS,Transformer
from shapely.geometry import shape,mapping
from shapely.ops import transform
from shapely import union_all
root=Path(__file__).resolve().parents[1];out=root/'public'/'data';out.mkdir(exist_ok=True,parents=True)
book=Path(sys.argv[1]) if len(sys.argv)>1 else Path(r'C:\Users\liu02034\Downloads\Uganda Ecosystem Services & PBIS - Hydroshed scoring - TEMP.xlsx')
archive=Path(sys.argv[2]) if len(sys.argv)>2 else Path(r'C:\Users\liu02034\Downloads\Catchments-20261016T002049Z-1-001.zip')
w=openpyxl.load_workbook(book,data_only=True,read_only=True)
raw=list(w['RawData'].values);headers=raw[0];records=[dict(zip(headers,r)) for r in raw[1:] if isinstance(r[19],(int,float))]
fields=['usle_per_ha','usle_Total','sed_export_per_ha','sed_export_Total','sed_deposition_per_ha','sed_deposition_Total','avoid_exp_per_ha','avoid_exp_Total','avoided_erosion_per_ha','avoided_erosion_Total','n_total_export_per_ha','n_total_export_Total','n_surface_export_per_ha','n_subsurface_export_per_ha','p_surface_export_per_ha','p_surface_export_Total','Deg_Share','Stable_Share','Improving_Share','wyield_mn','wyield_vol']
percentcols=['H','AI','K','AL','N','AO','Q','AR','T','AU','AC','BD','Z','W','AF','BG','BJ','BM','BP','BS','BV']
definitions=['Potential annual sheet and rill soil loss; baseline erosion pressure.','Total modeled annual sheet and rill soil loss.','Modeled sediment reaching streams per unit area; baseline pressure.','Total modeled annual sediment delivered to streams.','Modeled sediment deposition per unit area.','Total modeled annual sediment deposition.','Modeled sediment export avoided by current land cover; retention benefit, not a forecast of restoration gains.','Total modeled avoided sediment export.','Modeled soil erosion avoided by current land cover.','Total modeled avoided soil erosion.','Total nitrogen export per unit area.','Total nitrogen export load.','Surface nitrogen export per unit area.','Subsurface nitrogen export per unit area.','Surface phosphorus export per unit area.','Total surface phosphorus export load.','Fraction of mapped land classified as degraded.','Fraction of mapped land classified as stable.','Fraction of mapped land classified as improving.','Modeled annual water yield depth; not seasonal baseflow.','Modeled total annual water yield volume.']
indicators=[]
for i,row in enumerate(w['FinalRank'].iter_rows(min_row=9,max_row=29,min_col=1,max_col=6,values_only=True)):
 pillar,name,include,direction,weight,norm=row
 indicators.append(dict(id=fields[i],name=name,pillar=pillar,include=bool(include),direction='lower' if direction.startswith('Lower') else 'higher',weight=weight,unit=name[name.rfind('(')+1:-1],definition=definitions[i],percentileColumn=percentcols[i]))
z=zipfile.ZipFile(archive)
def read_shape(stem):
 base='Catchments/'+stem
 r=shapefile.Reader(shp=__import__('io').BytesIO(z.read(base+'.shp')),shx=__import__('io').BytesIO(z.read(base+'.shx')),dbf=__import__('io').BytesIO(z.read(base+'.dbf')),encoding='utf-8')
 crs=CRS.from_wkt(z.read(base+'.prj').decode());tr=Transformer.from_crs(crs,4326,always_xy=True)
 features=[];invalid=0
 for sr in r.iterShapeRecords():
  g=shape(sr.shape.__geo_interface__)
  if not g.is_valid:invalid+=1;g=g.buffer(0)
  g=transform(tr.transform,g)
  # Full-detail geometry retained; compact coordinates, no topology-changing simplification.
  features.append({'type':'Feature','properties':sr.record.as_dict(),'geometry':mapping(g)})
 return {'type':'FeatureCollection','features':features},crs.to_string(),invalid
geo,crs,invalid=read_shape('Uganda_WWFLev12_v01_S01');wmz,wcrs,winvalid=read_shape('Uganda_WMZs_v01')
print('Shape sample',geo['features'][0]['properties']);print('WMZ sample',wmz['features'][0]['properties'])
rawmap={int(r['Lev12_ID']):r for r in records};spmap={int(f['properties']['Lev12_ID']):f for f in geo['features']}
duplicates=lambda ids:[{'id':k,'count':v} for k,v in collections.Counter(ids).items() if v>1]
cross=[]
for key in ['Lev12_ID','ws_id','HYBAS_ID']:
 a=[r[key] for r in records if r.get(key) is not None];b=[f['properties'].get(key) for f in geo['features'] if f['properties'].get(key) is not None]
 cross.append({'field':key,'workbookDistinct':len(set(a)),'spatialDistinct':len(set(b)),'sharedDistinct':len(set(a)&set(b)),'workbookDuplicates':duplicates(a),'spatialDuplicates':duplicates(b)})
mismatches=[]
for id in rawmap.keys()&spmap.keys():
 a=rawmap[id];b=spmap[id]['properties']
 for field in ['HYBAS_ID','ws_id']:
  if a.get(field)!=b.get(field):mismatches.append({'Lev12_ID':id,'field':field,'workbook':a.get(field),'spatial':b.get(field)})
report={'workbook':book.name,'archive':archive.name,'workbookSHA256':hashlib.sha256(book.read_bytes()).hexdigest(),'archiveSHA256':hashlib.sha256(archive.read_bytes()).hexdigest(),'rawRecords':len(records),'spatialRecords':len(geo['features']),'matched':len(rawmap.keys()&spmap.keys()),'workbookOnly':sorted(rawmap.keys()-spmap.keys()),'spatialOnly':sorted(spmap.keys()-rawmap.keys()),'joinKey':'Lev12_ID','candidateKeys':cross,'crossAttributeMismatches':mismatches,'crs':crs,'invalidInputGeometriesRepaired':invalid,'wmzInvalidGeometriesRepaired':winvalid}
base=[r for r in records if r['AreaHa']>=1000]
def percentile_map(rows,field):
 values=sorted(float(r[field] or 0) for r in rows)
 return {int(r['Lev12_ID']):math.floor((bisect.bisect_left(values,float(r[field] or 0))/(len(values)-1))*1000+1e-10)/1000 if len(values)>1 else None for r in rows}
pcts={field:percentile_map(base,field) for field in fields};den=sum(i['weight'] for i in indicators if i['include'])
cached={int(r[0]):r for r in w['RankScoreData'].iter_rows(min_row=7,values_only=True) if isinstance(r[0],(int,float))}
scores={int(r['Lev12_ID']):sum(i['weight']/den*((1-pcts[i['id']][int(r['Lev12_ID'])]) if i['direction']=='lower' else pcts[i['id']][int(r['Lev12_ID'])]) for i in indicators if i['include']) for r in base}
errors=[abs(score-cached[id][74]) for id,score in scores.items()];perrors=[]
for i in indicators:
 col=openpyxl.utils.column_index_from_string(i['percentileColumn'])-1
 perrors.extend(abs(pcts[i['id']][id]-cached[id][col]) for id in scores)
report.update(defaultIncluded=len(base),defaultActiveIndicators=sum(i['include'] for i in indicators),defaultPositiveWeightIndicators=sum(i['include'] and i['weight']>0 for i in indicators),defaultRawWeightSum=den,validatedScores=len(errors),maxScoreDifference=max(errors),maxPercentileDifference=max(perrors),topFive=[{'Lev12_ID':id,'score':scores[id]} for id in sorted(scores,key=lambda id:(-scores[id],id))[:5]],missingValues={f:sum(r.get(f) is None for r in records) for f in fields},missingDefaultValues={f:sum(r.get(f) is None for r in base) for f in fields},defaultMissingAdministrativeLabels=sum(not r.get('Water_Mana') for r in base))
rows=[]
for r in records:
 id=int(r['Lev12_ID']);rows.append({'id':id,'hybasId':r['HYBAS_ID'],'wsId':r['ws_id'],'wmz':r.get('Water_Mana') or 'Unassigned','basin':r.get('Basin') or 'Unassigned','catchment':r.get('Catchment') or 'Unassigned','area':r['AreaHa'],'hasGeometry':id in spmap,'raw':{f:r.get(f) for f in fields},'sourceAttributes':r,'excelScore':cached[id][74] if id in cached else None})
def write(name,data):
 (out/name).write_text(json.dumps(data,ensure_ascii=False,separators=(',',':'),allow_nan=False),encoding='utf-8')
groups=collections.defaultdict(list)
for f in wmz['features']:groups[f['properties']['Water_Mana']].append(f)
dissolved={'type':'FeatureCollection','features':[{'type':'Feature','properties':{'Water_Mana':name,'WMZ_ID':fs[0]['properties']['WMZ_ID'],'sourceFeatureCount':len(fs)},'geometry':mapping(union_all([shape(f['geometry']) for f in fs]))} for name,fs in groups.items()]}
report['wmzSourceFeatures']=len(wmz['features']);report['wmzDissolvedFeatures']=len(dissolved['features'])
write('watersheds.json',rows);write('indicators.json',indicators);write('watersheds.geojson',geo);write('wmz-source.geojson',wmz);write('wmz.geojson',dissolved);write('validation.json',report)
(root/'reports'/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:report[k] for k in ['rawRecords','spatialRecords','matched','workbookOnly','spatialOnly','maxScoreDifference','maxPercentileDifference']},indent=2))
